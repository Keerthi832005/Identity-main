"""Builds IAM bulk-import workbooks from the legacy DMS_V2 database.

Output matches the contract in doc/IAM-Bulk-Data.md: sheets Data / Reference / _meta,
instruction band on row 1, headers on row 2, data from row 3.
"""

from __future__ import annotations

import argparse
import os
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SQLCMD = r"C:\Program Files\Microsoft SQL Server\Client SDK\ODBC\170\Tools\Binn\SQLCMD.EXE"
SEP = "\x1f"

HEADER_FILL = PatternFill("solid", fgColor="F4F4F6")
INSTRUCTION_FILL = PatternFill("solid", fgColor="FBE9EB")
REQUIRED_FONT = Font(bold=True, color="D01126")
HEADER_FONT = Font(bold=True, color="1F2126")
BOTTOM_BORDER = Border(bottom=Side(style="thin"))

# (columnId, header, required, width, helpText, allowedValuesKey)
ORG_COLUMNS = [
    ("unitType", "Unit type", True, 16,
     "One of the eight approved types. The parent's type decides what is allowed here.", "unitType"),
    ("unitCode", "Unit code", True, 18,
     "Unique across the whole organization, including other unit types.", None),
    ("unitName", "Unit name", True, 28, None, None),
    ("parentUnitCode", "Parent unit code", False, 18,
     "Blank only for an Organization row. A parent may be created by another row in this file.", None),
    ("description", "Description", False, 30, None, None),
    ("addressLine1", "Address line 1", False, 26, None, None),
    ("addressLine2", "Address line 2", False, 26, None, None),
    ("addressLine3", "Address line 3", False, 26, None, None),
    ("city", "City", False, 18, None, None),
    ("district", "District", False, 18, None, None),
    ("stateName", "State", False, 18, None, None),
    ("postalCode", "Postal code", False, 14, None, None),
    ("countryCode", "Country code", False, 14, None, None),
    ("latitude", "Latitude", False, 14, None, None),
    ("longitude", "Longitude", False, 14, None, None),
]
ORG_INSTRUCTION = (
    "One unit per row. A parent may appear anywhere in this file; order does not matter. "
    "Leave the parent blank only on an Organization row."
)
ORG_TYPES = ["Organization", "Country", "Region", "State", "Branch", "Location", "Department", "Team"]

USER_COLUMNS = [
    ("employeeCode", "Employee code", True, 18, "The code the person signs in with. Must be unique.", None),
    ("displayName", "Display name", True, 28, "Shown in the directory and on every audit event.", None),
    ("email", "Contact email", False, 32,
     "Optional. Contact only - it does not enable email sign-in or recovery.", None),
    ("managerEmployeeCode", "Manager employee code", False, 22,
     "The manager's employee code. They may be created earlier in this same file.", "managers"),
]
USER_INSTRUCTION = (
    "One user per row. Employee code and display name are required. "
    "Passwords, PINs and MFA are never set here."
)

MAX_LEN = {
    "unitCode": 50, "unitName": 200, "parentUnitCode": 50, "description": 500,
    "addressLine1": 250, "addressLine2": 250, "addressLine3": 250,
    "city": 100, "district": 100, "stateName": 100, "postalCode": 20, "countryCode": 3,
    "employeeCode": 50, "displayName": 200, "email": 256, "managerEmployeeCode": 50,
}


def query(server: str, database: str, sql: str) -> list[list[str]]:
    """Runs one query and returns rows of trimmed text. SEP cannot occur in the source data."""
    with tempfile.TemporaryDirectory() as folder:
        out = Path(folder) / "out.txt"
        result = subprocess.run(
            [SQLCMD, "-S", server, "-d", database, "-E", "-C", "-h", "-1", "-W",
             "-s", SEP, "-Q", "SET NOCOUNT ON; " + sql, "-o", str(out)],
            capture_output=True, text=True,
        )
        text = out.read_text(encoding="utf-8", errors="replace") if out.exists() else ""
    if result.returncode != 0:
        raise SystemExit(f"sqlcmd failed ({result.returncode}):\n{result.stderr}\n{text}")
    rows = []
    for line in text.splitlines():
        if not line.strip():
            continue
        if line.startswith("Msg ") or line.startswith("Sqlcmd:"):
            raise SystemExit(f"sqlcmd reported an error:\n{text}")
        rows.append([("" if cell == "NULL" else cell).strip() for cell in line.split(SEP)])
    return rows


def slug(value: str, limit: int = 44) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9]+", "-", value).strip("-").upper()
    return cleaned[:limit].strip("-")


def clip(column_id: str, value: str) -> str:
    limit = MAX_LEN.get(column_id)
    value = " ".join(value.split())
    return value[:limit] if limit else value


def write_sheet(path: Path, entity_key: str, version: int, columns, instruction, rows, references):
    book = Workbook()
    data = book.active
    data.title = "Data"

    last = get_column_letter(len(columns))
    data.merge_cells(f"A1:{last}1")
    band = data["A1"]
    band.value = instruction
    band.fill = INSTRUCTION_FILL
    band.alignment = Alignment(wrap_text=True, vertical="center")
    data.row_dimensions[1].height = 34

    for index, (column_id, header, required, width, help_text, _) in enumerate(columns, start=1):
        cell = data.cell(row=2, column=index)
        cell.value = header + " *" if required else header
        cell.font = REQUIRED_FONT if required else HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BOTTOM_BORDER
        if help_text:
            cell.comment = Comment(help_text, entity_key)
        data.column_dimensions[get_column_letter(index)].width = width
    data.freeze_panes = "A3"

    for offset, row in enumerate(rows):
        for index, (column_id, *_rest) in enumerate(columns, start=1):
            value = row.get(column_id)
            if value in (None, ""):
                continue
            cell = data.cell(row=3 + offset, column=index)
            cell.value = value
            cell.number_format = "@"

    reference = book.create_sheet("Reference")
    for index, (key, values) in enumerate(references, start=1):
        letter = get_column_letter(index)
        head = reference.cell(row=1, column=index)
        head.value = key
        head.font = Font(bold=True)
        for offset, value in enumerate(values):
            reference.cell(row=2 + offset, column=index).value = value
        reference.column_dimensions[letter].width = 28
        target = next((i for i, c in enumerate(columns, start=1) if c[5] == key), None)
        if target and values:
            rule = DataValidation(
                type="list",
                formula1=f"=Reference!${letter}$2:${letter}${len(values) + 1}",
                allow_blank=True,
            )
            rule.showInputMessage = False
            data.add_data_validation(rule)
            column_letter = get_column_letter(target)
            rule.add(f"{column_letter}3:{column_letter}{max(len(rows) + 2, 3)}")

    meta = book.create_sheet("_meta")
    meta["A1"], meta["B1"] = "entityKey", entity_key
    meta["A2"], meta["B2"] = "templateVersion", str(version)
    meta["A3"], meta["B3"] = "columnIds", ",".join(c[0] for c in columns)
    meta.sheet_state = "hidden"

    book.save(path)
    return len(rows)


def build_organization_rows(server: str, database: str, org_code: str, org_name: str):
    country = query(server, database, (
        "SELECT alpha_3code, country_name FROM countries "
        "WHERE id = (SELECT TOP 1 country_id FROM branch_locations WHERE country_id IS NOT NULL);"
    ))
    country_code, country_name = country[0][0], country[0][1]

    sites = query(server, database, (
        "SELECT b.id, b.erp_organization_id, b.location_name,"
        " ISNULL(b.city,''), ISNULL(b.state,''), ISNULL(b.region,''), ISNULL(b.postal_code,'')"
        " FROM branch_locations b WHERE b.available = 1 ORDER BY b.id;"
    ))
    departments = query(server, database,
                        "SELECT id, main_department_name FROM main_departments WHERE available = 1 ORDER BY id;")
    teams = query(server, database, (
        "SELECT s.id, s.sub_department_name, d.main_department_name FROM sub_departments s"
        " JOIN main_departments d ON d.id = s.main_department_id"
        " WHERE s.available = 1 AND d.available = 1 ORDER BY s.main_department_id, s.id;"
    ))

    rows = [{"unitType": "Organization", "unitCode": org_code, "unitName": org_name,
             "description": "Root organization migrated from DMS_V2."}]
    rows.append({"unitType": "Country", "unitCode": country_code, "unitName": country_name,
                 "parentUnitCode": org_code, "countryCode": country_code})

    regions, states, branches = {}, {}, {}
    # branch_locations.address holds the region word, not a street address, so it is not mapped.
    for _id, erp, name, city, state, region, postal in sites:
        region = region or "Unassigned"
        state = state or "Unassigned"
        city = city or state

        region_code = "RGN-" + slug(region)
        if region_code not in regions:
            regions[region_code] = True
            rows.append({"unitType": "Region", "unitCode": region_code, "unitName": region,
                         "parentUnitCode": country_code, "countryCode": country_code})

        state_code = "ST-" + slug(state)
        if state_code not in states:
            states[state_code] = region_code
            rows.append({"unitType": "State", "unitCode": state_code, "unitName": state,
                         "parentUnitCode": region_code, "stateName": state, "countryCode": country_code})

        branch_code = "BR-" + slug(city)
        if branch_code not in branches:
            branches[branch_code] = state_code
            rows.append({"unitType": "Branch", "unitCode": branch_code, "unitName": city,
                         "parentUnitCode": state_code, "city": city, "stateName": state,
                         "countryCode": country_code})

        rows.append({
            "unitType": "Location", "unitCode": "LOC-" + slug(name), "unitName": name,
            "parentUnitCode": branch_code,
            "description": f"DMS branch location {_id}; ERP organization {erp}.",
            "city": city, "stateName": state,
            "postalCode": postal, "countryCode": country_code,
        })

    for _id, name in departments:
        rows.append({"unitType": "Department", "unitCode": "DEP-" + slug(name), "unitName": name,
                     "parentUnitCode": org_code})

    department_codes = {name: "DEP-" + slug(name) for _id, name in departments}
    for _id, name, parent_name in teams:
        rows.append({"unitType": "Team", "unitCode": "TM-" + slug(name), "unitName": name,
                     "parentUnitCode": department_codes[parent_name]})

    return [{k: clip(k, v) for k, v in row.items()} for row in rows]


def build_user_rows(server: str, database: str, prefix: str):
    people = query(server, database, (
        "SELECT e.employee_number, e.employee_name, ISNULL(e.email,''), ISNULL(m.employee_number,''),"
        " e.active_or_inactive"
        " FROM employees e LEFT JOIN employees m ON m.id = e.manager_id"
        " WHERE e.employee_number IS NOT NULL AND LTRIM(RTRIM(e.employee_number)) <> ''"
        " ORDER BY e.id;"
    ))
    # IAM signs people in as INDE<employee_number>; the bare DMS number is not an IAM identity.
    codes = {prefix + row[0] for row in people}
    rows, dropped = [], 0
    for number, name, email, manager_number, _active in people:
        code = prefix + number
        manager = prefix + manager_number if manager_number else ""
        if manager and (manager == code or manager not in codes):
            manager, dropped = "", dropped + 1
        rows.append({k: clip(k, v) for k, v in {
            "employeeCode": code, "displayName": name, "email": email,
            "managerEmployeeCode": manager,
        }.items()})
    inactive = sum(1 for row in people if row[4] == "0")
    return rows, dropped, inactive


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--server", default=os.environ.get("DMS_SERVER", r"HOCOM18502627\SQLEXPRESS2022"))
    parser.add_argument("--database", default=os.environ.get("DMS_DATABASE", "DMS_V2"))
    parser.add_argument("--org-code", default="FIPL")
    parser.add_argument("--org-name", default="Fujitec India Pvt Ltd")
    parser.add_argument("--code-prefix", default="INDE",
                        help="Prefixed to the DMS employee number to form the IAM employee code.")
    parser.add_argument("--out", default=str(Path(__file__).parent))
    args = parser.parse_args()

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    org_rows = build_organization_rows(args.server, args.database, args.org_code, args.org_name)
    duplicates = {r["unitCode"] for r in org_rows if
                  sum(1 for x in org_rows if x["unitCode"] == r["unitCode"]) > 1}
    if duplicates:
        raise SystemExit(f"Unit codes are not unique: {sorted(duplicates)}")

    write_sheet(out / "organization-units.xlsx", "organization-units", 1, ORG_COLUMNS,
                ORG_INSTRUCTION, org_rows, [("unitType", ORG_TYPES)])

    user_rows, dropped, inactive = build_user_rows(args.server, args.database, args.code_prefix)
    managers = sorted({r["managerEmployeeCode"] for r in user_rows if r["managerEmployeeCode"]})
    write_sheet(out / "users.xlsx", "users", 1, USER_COLUMNS,
                USER_INSTRUCTION, user_rows, [("managers", managers)])

    counts = {}
    for row in org_rows:
        counts[row["unitType"]] = counts.get(row["unitType"], 0) + 1
    print(f"organization-units.xlsx  {len(org_rows)} rows  " +
          " ".join(f"{k}={counts[k]}" for k in ORG_TYPES if k in counts))
    print(f"users.xlsx               {len(user_rows)} rows  "
          f"managers={len(managers)} manager-links-dropped={dropped} inactive-in-dms={inactive}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
