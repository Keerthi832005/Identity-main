"""Re-runs the server's staging rules against the generated workbooks before anyone uploads them.

Mirrors ExcelWorkbookReader plus OrganizationBulkValidator / UsersBulkValidator, so a failure here
is the same failure the IMPORT screen would report.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import load_workbook

from build_bulk_workbooks import MAX_LEN, ORG_COLUMNS, ORG_TYPES, USER_COLUMNS

PARENT_TYPE = {
    "Organization": None, "Country": "Organization", "Region": "Country", "State": "Region",
    "Branch": "State", "Location": "Branch", "Department": "Organization", "Team": "Department",
}


def normalise(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9]", "", (text or "").strip().rstrip("*").strip()).lower()


def read(path: Path, entity_key: str, columns):
    book = load_workbook(path, data_only=True)
    meta = book["_meta"]
    assert meta["B1"].value == entity_key, f"{path.name}: entityKey is {meta['B1'].value!r}"
    assert str(meta["B2"].value) == "1", f"{path.name}: templateVersion is {meta['B2'].value!r}"
    assert meta.sheet_state == "hidden", f"{path.name}: _meta is visible"
    assert "Reference" in book.sheetnames, f"{path.name}: no Reference sheet"

    data = book["Data"]
    headers = {}
    for cell in data[2]:
        key = normalise(cell.value)
        if key and key not in headers:
            headers[key] = cell.column
    positions = {}
    for column_id, header, required, *_ in columns:
        index = headers.get(normalise(header)) or headers.get(normalise(column_id))
        assert index or not required, f"{path.name}: required column {header!r} missing"
        if index:
            positions[column_id] = index

    rows = []
    for excel_row in range(3, data.max_row + 1):
        values = {}
        for column_id, index in positions.items():
            value = data.cell(row=excel_row, column=index).value
            if value is not None and str(value).strip():
                values[column_id] = str(value).strip()
        if values:
            rows.append((excel_row, values))
    return rows


def check_cells(rows, columns, key_ids, errors):
    seen = {}
    for excel_row, values in rows:
        for column_id, header, required, _w, _h, allowed_key in columns:
            value = values.get(column_id)
            if required and not value:
                errors.append(f"row {excel_row}: cell.required on {header}")
            if value and column_id in MAX_LEN and len(value) > MAX_LEN[column_id]:
                errors.append(f"row {excel_row}: cell.too-long on {header} ({len(value)})")
        key = tuple(values.get(k, "").lower() for k in key_ids)
        if all(key):
            if key in seen:
                errors.append(f"row {excel_row}: row.duplicate-in-batch with row {seen[key]}")
            seen[key] = excel_row


def check_organizations(rows, errors):
    types = {v["unitCode"].lower(): v.get("unitType") for _r, v in rows if v.get("unitCode")}
    parents = {v["unitCode"].lower(): v.get("parentUnitCode", "").lower()
               for _r, v in rows if v.get("unitCode")}
    for excel_row, values in rows:
        unit_type = values.get("unitType")
        code = values.get("unitCode", "").lower()
        parent = values.get("parentUnitCode", "").lower()
        if unit_type not in ORG_TYPES:
            errors.append(f"row {excel_row}: unit.type-unknown {unit_type!r}")
            continue
        expected = PARENT_TYPE[unit_type]
        if expected is None:
            if parent:
                errors.append(f"row {excel_row}: unit.root-has-parent")
            continue
        if not parent:
            errors.append(f"row {excel_row}: unit.parent-required ({unit_type} needs a {expected})")
        elif parent == code:
            errors.append(f"row {excel_row}: unit.parent-self")
        elif parent not in types:
            errors.append(f"row {excel_row}: unit.parent-unknown {parent}")
        elif types[parent] != expected:
            errors.append(f"row {excel_row}: unit.parent-wrong-type ({parent} is {types[parent]}, "
                          f"expected {expected})")
        walk, seen = parent, set()
        while walk and walk in parents:
            if walk in seen or walk == code:
                errors.append(f"row {excel_row}: unit.parent-cycle")
                break
            seen.add(walk)
            walk = parents[walk]


def plausible_email(value: str) -> bool:
    at = value.find("@")
    if at <= 0 or at != value.rfind("@"):
        return False
    domain = value[at + 1:]
    return (len(domain) >= 3 and "." in domain and not domain.startswith(".")
            and not domain.endswith(".") and " " not in value)


def check_users(rows, errors):
    codes = {v["employeeCode"].lower() for _r, v in rows if v.get("employeeCode")}
    for excel_row, values in rows:
        code = values.get("employeeCode", "").lower()
        manager = values.get("managerEmployeeCode", "").lower()
        email = values.get("email")
        if manager:
            if manager == code:
                errors.append(f"row {excel_row}: manager.self")
            elif manager not in codes:
                errors.append(f"row {excel_row}: manager.unknown {manager}")
        if email and not plausible_email(email):
            errors.append(f"row {excel_row}: email.invalid {email!r}")


def main() -> int:
    folder = Path(__file__).parent
    failures = 0

    org_rows = read(folder / "organization-units.xlsx", "organization-units", ORG_COLUMNS)
    errors: list[str] = []
    check_cells(org_rows, ORG_COLUMNS, ["unitCode"], errors)
    check_organizations(org_rows, errors)
    print(f"organization-units.xlsx  {len(org_rows)} rows  {len(errors)} errors")
    for line in errors[:20]:
        print("  " + line)
    failures += len(errors)

    user_rows = read(folder / "users.xlsx", "users", USER_COLUMNS)
    errors = []
    check_cells(user_rows, USER_COLUMNS, ["employeeCode"], errors)
    check_users(user_rows, errors)
    print(f"users.xlsx               {len(user_rows)} rows  {len(errors)} errors")
    for line in errors[:20]:
        print("  " + line)
    failures += len(errors)

    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
